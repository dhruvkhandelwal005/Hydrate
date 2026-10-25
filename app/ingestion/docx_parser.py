from __future__ import annotations
from pathlib import Path
from typing import Tuple
import docx

def parse_docx_file(path: Path) -> Tuple[str, list[list[str]], str | None]:
    """
    Parses a DOCX file into full text with preserved paragraph tags [P1], [P2],
    plus extracted tables as rows of strings.
    """
    try:
        doc = docx.Document(path)
    except Exception as e:
        return "", [], f"DOCX open error: {type(e).__name__}: {e}"

    lines: list[str] = []
    p_num = 1
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        # Preserve existing [P1] locators if present, otherwise assign locator
        if not txt.startswith("[P"):
            lines.append(f"[P{p_num}] {txt}")
        else:
            lines.append(txt)
        p_num += 1

    extracted_tables: list[list[str]] = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells]
            if any(row_cells):
                extracted_tables.append(row_cells)
                lines.append(" | ".join(row_cells))

    full_text = "\n\n".join(lines)
    return full_text, extracted_tables, None
